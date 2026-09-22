"""Backfills historical eval runs into Langfuse as Experiment runs against the
'eu-ai-act-golden-set' Dataset, so the Compare Experiments view shows our actual
project history. Replays already-computed scores/answers from files we already
saved — makes zero new LLM calls (no Groq, no judge), per DECISIONS.md.

Run with: python -m eval.langfuse_backfill_experiments
"""

import json
from pathlib import Path

import openpyxl
from langfuse import Langfuse

from app.core import config
from eval.pipeline_metadata import build_pipeline_metadata

DATASET_NAME = "eu-ai-act-golden-set"
EVAL_DIR = Path(__file__).parent


def _load_score_records(path: Path) -> dict[str, dict]:
    """dense/hybrid/hybrid_rerank raw-results files: {question: {"scores": {...}, "retrieval_hit": bool?}}"""
    records = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        scores = dict(rec["scores"])
        if "retrieval_hit" in rec:
            scores["Retrieval Hit"] = 1.0 if rec["retrieval_hit"] else 0.0
        records[rec["question"]] = {"output": None, "scores": scores}
    return records


def _load_claude_judge_records() -> dict[str, dict]:
    """Combines pipeline_outputs_for_external_judge.jsonl (real generated answers)
    with rag_eval_report.xlsx's Per-Item Scores sheet (Claude's judge scores)."""
    outputs = {}
    for line in (EVAL_DIR / "pipeline_outputs_for_external_judge.jsonl").read_text(
        encoding="utf-8"
    ).splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        outputs[rec["question"]] = rec["generated_answer"]

    wb = openpyxl.load_workbook(EVAL_DIR / "rag_eval_report.xlsx", data_only=True)
    ws = wb["Per-Item Scores"]
    rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))

    records = {}
    for row in rows:
        if row[0] is None or row[2] is None:  # blank/total row
            continue
        question = row[2]
        scores = {
            "Retrieval Hit": float(row[5]) if row[5] is not None else None,
            "MRR": row[7],
            "Faithfulness": row[10],
            "Answer Correctness": row[11],
        }
        scores = {k: v for k, v in scores.items() if v is not None}
        records[question] = {
            "output": outputs.get(question),
            "scores": scores,
            "comment": row[12],  # Outcome Label
        }
    return records


def _load_contextual_headers_full_judge_records() -> dict[str, dict]:
    """Combines contextual_headers_full_outputs_for_external_judge.jsonl (real
    generated answers, contextual-headers pipeline) with
    contextual_headers_full_judge_scores.jsonl (Claude Sonnet 5's judge
    scores) for all 41 golden-set questions (Step 10 -- extends Step 8's
    15-question hard subset to the full set)."""
    outputs = {}
    for line in (EVAL_DIR / "contextual_headers_full_outputs_for_external_judge.jsonl").read_text(
        encoding="utf-8"
    ).splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        outputs[rec["question"]] = rec["generated_answer"]

    records = {}
    for line in (EVAL_DIR / "contextual_headers_full_judge_scores.jsonl").read_text(
        encoding="utf-8"
    ).splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        question = rec["question"]
        scores = {
            "Retrieval Hit": 1.0 if rec["retrieval_hit"] else 0.0,
            "Chunk-Level Hit": 1.0 if rec["chunk_level_hit"] else 0.0,
            "MRR": rec["mrr"],
            "Faithfulness": rec["faithfulness"],
            "Answer Correctness": rec["answer_correctness"],
            "Citation Accuracy": rec["citation_accuracy"],
        }
        records[question] = {
            "output": outputs.get(question),
            "scores": scores,
            "comment": f"{rec['outcome_label']} — {rec['comment']}",
        }
    return records


def _load_hard_subset_judge_records() -> dict[str, dict]:
    """Combines hard_subset_outputs_for_external_judge.jsonl (real generated
    answers, contextual-headers pipeline) with hard_subset_judge_scores.jsonl
    (Claude Sonnet 5's judge scores) for the 15-question hard subset only —
    the other 26 golden-set questions get no score/output for this run."""
    outputs = {}
    for line in (EVAL_DIR / "hard_subset_outputs_for_external_judge.jsonl").read_text(
        encoding="utf-8"
    ).splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        outputs[rec["question"]] = rec["generated_answer"]

    records = {}
    for line in (EVAL_DIR / "hard_subset_judge_scores.jsonl").read_text(
        encoding="utf-8"
    ).splitlines():
        if not line.strip():
            continue
        rec = json.loads(line)
        question = rec["question"]
        scores = {
            "Retrieval Hit": 1.0 if rec["retrieval_hit"] else 0.0,
            "Chunk-Level Hit": 1.0 if rec["chunk_level_hit"] else 0.0,
            "MRR": rec["mrr"],
            "Faithfulness": rec["faithfulness"],
            "Answer Correctness": rec["answer_correctness"],
            "Citation Accuracy": rec["citation_accuracy"],
        }
        records[question] = {
            "output": outputs.get(question),
            "scores": scores,
            "comment": f"{rec['outcome_label']} — {rec['comment']}",
        }
    return records


def backfill_run(
    client: Langfuse,
    dataset,
    run_name: str,
    description: str,
    records: dict[str, dict],
    metadata_overrides: dict,
) -> None:
    def task(*, item, **kwargs):
        rec = records.get(item.input)
        return rec["output"] if rec else None

    def evaluator(*, input, output, expected_output, metadata, **kwargs):
        rec = records.get(input)
        if not rec:
            return []
        return [
            {"name": name, "value": value, "comment": rec.get("comment")}
            for name, value in rec["scores"].items()
        ]

    result = dataset.run_experiment(
        name=run_name,
        run_name=run_name,
        description=description,
        task=task,
        evaluators=[evaluator],
        metadata=build_pipeline_metadata(metadata_overrides),
    )
    print(f"  -> {run_name}: {len(result.item_results)} items")


def main() -> None:
    if not (config.LANGFUSE_PUBLIC_KEY and config.LANGFUSE_SECRET_KEY):
        raise RuntimeError("LANGFUSE_PUBLIC_KEY/LANGFUSE_SECRET_KEY not set in .env")

    client = Langfuse(
        public_key=config.LANGFUSE_PUBLIC_KEY,
        secret_key=config.LANGFUSE_SECRET_KEY,
        host=config.LANGFUSE_HOST,
    )
    dataset = client.get_dataset(DATASET_NAME)

    runs = [
        (
            "phase2-dense-only-fixed500-baseline-27of41",
            "Phase 2 baseline: dense-only retrieval, fixed 500/50 chunking, post-bugfix. 27/41 (66%) on 4 pipeline metrics.",
            _load_score_records(EVAL_DIR / "eval_run_raw_results_dense.jsonl"),
            {"retrieval_mode": "dense", "qdrant_collection": "ai_act_corpus"},
        ),
        (
            "phase3-hybrid-fixed500-reranked-REVERTED-26of41",
            "Phase 3: hybrid search + cross-encoder reranking. REVERTED - pass rate dropped to 26/41 despite 2 of 4 metrics improving on average.",
            _load_score_records(EVAL_DIR / "eval_run_raw_results_hybrid_rerank.jsonl"),
            {"retrieval_mode": "hybrid", "use_reranking": True, "qdrant_collection": "ai_act_corpus_hybrid"},
        ),
        (
            "phase3-hybrid-fixed500-reference-24of41",
            "Phase 3 current reference point: hybrid search (kept), fixed 500/50 chunking, all 6 checks (4 pipeline + retrieval-hit + answer-correctness). 24/41 (59%).",
            _load_score_records(EVAL_DIR / "eval_run_raw_results_hybrid.jsonl"),
            {"retrieval_mode": "hybrid", "qdrant_collection": "ai_act_corpus_hybrid"},
        ),
        (
            "phase3-hybrid-recursive500-sonnet5judged-34of41",
            "Phase 3: hybrid search + recursive chunking, judged independently by Claude Sonnet 5 (not Groq). 83% doc-level retrieval recall - confirms the Groq-free screen.",
            _load_claude_judge_records(),
            {"retrieval_mode": "hybrid", "chunking_strategy": "recursive", "qdrant_collection": "ai_act_corpus_recursive_hybrid"},
        ),
        (
            "phase3-hybrid-fixed500-ctxheaders-sonnet5judged-hard15-10of15",
            "Phase 3 Step 8: hybrid search + contextual chunk headers, judged by Claude Sonnet 5 on a targeted 15-question hard subset (cross-reference/multi-hop/temporal-conflict + known retrieval misses), not the full 41. 67% doc-recall, 53% fully/mostly correct on this hardest-question subset -- not comparable to other runs' full-set percentages.",
            _load_hard_subset_judge_records(),
            {"retrieval_mode": "hybrid", "use_contextual_headers": True, "qdrant_collection": "ai_act_corpus_hybrid"},
        ),
        (
            "phase3-hybrid-fixed500-ctxheaders-sonnet5judged-full-34of41",
            "Phase 3 Step 10: hybrid search + contextual chunk headers, judged by Claude Sonnet 5 on all 41 questions (extends Step 8's 15-question hard subset). 36/41 (87.8%) doc-recall -- exact match with the Groq-free screen -- and 34/41 (82.9%) fully/mostly correct, the best full-set outcome-distribution result yet.",
            _load_contextual_headers_full_judge_records(),
            {"retrieval_mode": "hybrid", "use_contextual_headers": True, "qdrant_collection": "ai_act_corpus_hybrid"},
        ),
    ]

    print(f"Backfilling {len(runs)} historical experiment runs into dataset '{DATASET_NAME}':")
    for run_name, description, records, metadata_overrides in runs:
        backfill_run(client, dataset, run_name, description, records, metadata_overrides)

    client.flush()
    print("Done.")


if __name__ == "__main__":
    main()
